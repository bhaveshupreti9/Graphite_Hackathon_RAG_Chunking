"""Spreadsheet loading, summarised history and the /api/history endpoints."""

from __future__ import annotations

import datetime
import json
import time
from pathlib import Path

import openpyxl
import pytest
from fastapi.testclient import TestClient

from backend import memory, schemas
from baseline.naive import naive_chunk_document
from retrieval import llm_client
from smartchunk.chunker import chunk_document
from smartchunk.config import Settings
from smartchunk.damage import count_structural_damage
from smartchunk.errors import DocumentError
from smartchunk.pdf_loader import document_page_count, load_document, load_plain
from smartchunk.tabular import column_profile, to_number

# Spreadsheets -------------------------------------------------------------------------------------


def test_sample_csv_chunks_keep_tables_whole(samples_dir: Path) -> None:
    path = samples_dir / "sales_2026.csv"
    pages, warnings = load_document(path)
    text = pages[0].text
    assert "## Column overview" in text and "## Data" in text
    assert "| Revenue (INR) | number | 120 of 120 |" in text
    result = chunk_document(path, Settings(), "doc_0000aaaa")
    assert result.chunks and all(c.type in ("prose", "table") for c in result.chunks)
    damage = count_structural_damage(result.raw_blocks, result.chunks, result.oversized_blocks)
    assert damage["split_tables"] == 0
    for chunk in result.chunks:
        if chunk.type == "table" and chunk.heading_path[-1] == "Data":
            assert "| Row | Order ID |" in chunk.text  # every data chunk repeats the header row
    assert naive_chunk_document(path, Settings(), "doc_0000aaaa")
    assert document_page_count(path) == 1
    assert warnings == []


def test_excel_workbook(tmp_path: Path) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Staff"
    sheet.append(["Name", "Joined", "Salary", None, "Note"])
    sheet.append(["Asha", datetime.datetime(2024, 3, 1), 55000.0, None, "a|b"])
    sheet.append(["Ravi", datetime.datetime(2023, 7, 15), 72000.5, None, None])
    workbook.create_sheet("Empty")
    other = workbook.create_sheet("Stock")
    other.append(["Item", "Qty"])
    other.append(["Pen", 10])
    path = tmp_path / "book.xlsx"
    workbook.save(path)

    pages, _ = load_document(path)
    text = pages[0].text
    assert len(pages) == 1  # one page, so identical header rows are never treated as page headers
    assert "# Staff" in text and "# Stock" in text and "Empty" not in text
    assert "a\\|b" in text  # pipes inside cells are escaped
    assert "| Salary | number | 2 of 2 |" in text
    assert "2024-03-01" in text
    assert document_page_count(path) == 2
    result = chunk_document(path, Settings(), "doc_0000bbbb")
    assert {tuple(c.heading_path) for c in result.chunks} >= {("Staff", "Data"), ("Stock", "Data")}
    assert "Sheet: Staff" in load_plain(path)[0].text


def test_bad_spreadsheets(tmp_path: Path) -> None:
    empty = tmp_path / "empty.csv"
    empty.write_text(" \n,,\n", encoding="utf-8")
    with pytest.raises(DocumentError):
        load_document(empty)
    fake = tmp_path / "fake.xlsx"
    fake.write_bytes(b"this is not a zip file")
    with pytest.raises(DocumentError):
        load_document(fake)


def test_csv_variants(tmp_path: Path) -> None:
    semicolon = tmp_path / "eu.csv"
    semicolon.write_text("city;people\nPune;3000000\nGoa;1500000\n", encoding="utf-8")
    assert "| Pune | 3000000 |" in load_document(semicolon)[0][0].text
    tsv = tmp_path / "data.tsv"
    tsv.write_text("a\tb\n1\t2\n", encoding="utf-8")
    assert "| 1 | 1 | 2 |" in load_document(tsv)[0][0].text
    headless = tmp_path / "numbers.csv"
    headless.write_text("1,2\n3,4\n", encoding="utf-8")
    assert "| Row | Column 1 | Column 2 |" in load_document(headless)[0][0].text
    latin = tmp_path / "latin.csv"
    latin.write_bytes("name,city\nJosé,Málaga\n".encode("latin-1"))
    pages, warnings = load_document(latin)
    assert "José" in pages[0].text and warnings


def test_number_parsing_and_profile() -> None:
    assert to_number("₹1,250.50") == 1250.5
    assert to_number("(120)") == -120
    assert to_number("12%") == 12
    assert to_number("abc") is None and to_number("") is None
    profile = column_profile("Price", ["10", "20", "", "30"])
    assert profile["Type"] == "number" and profile["Sum"] == "60" and profile["Filled"] == "3 of 4"
    assert column_profile("City", ["Pune", "Pune", "Goa"])["Most common"] == "Pune (2)"


# Summarised history ------------------------------------------------------------------------------


@pytest.fixture
def sync_memory(monkeypatch: pytest.MonkeyPatch):
    monkeypatch.setattr(memory, "BACKGROUND", False)


def test_history_folds_into_summary_without_llm(sync_memory) -> None:
    for i in range(memory.RECENT_LIMIT + 3):
        memory.record("ana", "doc_11112222", "notes.md", f"Question {i}?", "smart", f"Answer {i}. More.", None, i == 0)
    data = memory.get("ana", "doc_11112222")
    assert len(data["recent"]) <= memory.RECENT_LIMIT
    assert data["summarized_questions"] == memory.COMPRESS_BATCH
    assert "Question 0? → not found in the document" in data["summary"]
    assert "Question 1? → Answer 1." in data["summary"]
    assert data["counts"]["questions"] == memory.RECENT_LIMIT + 3
    assert data["counts"]["not_found"] == 1
    schemas.DocHistory.model_validate(data)
    # Other users never see it.
    assert memory.get("bob", "doc_11112222")["counts"]["questions"] == 0


def test_history_folds_with_llm(sync_memory, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(llm_client, "is_configured", lambda: True)
    monkeypatch.setattr(llm_client, "generate", lambda *a, **k: ("They asked about loops and sorting.", None))
    for i in range(memory.RECENT_LIMIT + 1):
        memory.record("ana", "doc_33334444", "t.md", f"Q{i}", "compare", f"A{i}", f"N{i}", False)
    data = memory.get("ana", "doc_33334444")
    assert data["summary"] == "They asked about loops and sorting."
    assert data["recent"][0]["naive_answer"] == f"N{memory.COMPRESS_BATCH}"


def test_background_fold_finishes() -> None:
    for i in range(memory.RECENT_LIMIT + 1):
        memory.record("ana", "doc_55556666", "t.md", f"Q{i}", "naive", f"A{i}.", None, False)
    for _ in range(100):
        if memory.get("ana", "doc_55556666")["summarized_questions"]:
            break
        time.sleep(0.02)
    assert memory.get("ana", "doc_55556666")["summarized_questions"] == memory.COMPRESS_BATCH


def test_corrupt_history_file_is_ignored(tmp_path: Path) -> None:
    memory.record("ana", "doc_77778888", "t.md", "Q", "smart", "A", None, False)
    path = memory._file("ana", "doc_77778888")
    path.write_text("{not json", encoding="utf-8")
    assert memory.get("ana", "doc_77778888")["counts"]["questions"] == 0
    memory.record("ana", "doc_77778888", "t.md", "Q2", "smart", "A2", None, False)
    assert json.loads(path.read_text(encoding="utf-8"))["counts"]["questions"] == 1


# API ---------------------------------------------------------------------------------------------


@pytest.fixture
def client(monkeypatch: pytest.MonkeyPatch, tmp_path: Path):
    monkeypatch.setenv("SKIP_WARMUP", "1")
    monkeypatch.setattr(memory, "BACKGROUND", False)
    monkeypatch.setattr(llm_client, "is_configured", lambda: True)
    calls: list[list[dict[str, str]]] = []

    def generate(messages, temperature=0, max_tokens=512):
        calls.append(messages)
        if "Rewrite" in messages[0]["content"]:
            return "What is the total revenue?", None
        if messages[0]["role"] == "system":
            return json.dumps({"summary": "Sales data.", "key_topics": ["revenue"], "suggested_questions": ["Total?"]}), None
        return "| Metric | Value |\n|---|---|\n| Revenue | 27,378,282.50 |", llm_client.Usage(100, 20)

    monkeypatch.setattr(llm_client, "generate", generate)
    from backend.main import create_app

    with TestClient(create_app()) as test_client:
        token = test_client.post("/api/auth/register", json={"username": "maya", "password": "secret123"}).json()["token"]
        test_client.headers["Authorization"] = f"Bearer {token}"
        test_client.calls = calls  # type: ignore[attr-defined]
        yield test_client


def _process(client: TestClient, sample_id: str) -> str:
    doc_id = client.post("/api/upload-sample", json={"sample_id": sample_id}).json()["doc_id"]
    settings = Settings(not_found_threshold=0.0).to_dict()
    assert client.post("/api/process", json={"doc_id": doc_id, "settings": settings}).status_code == 202
    for _ in range(300):
        status = client.get(f"/api/processing-status/{doc_id}").json()
        if status["status"] in ("done", "error"):
            break
        time.sleep(0.05)
    assert status["status"] == "done", status
    return doc_id


def test_csv_upload_chat_and_history_endpoints(client: TestClient) -> None:
    samples = {s["sample_id"] for s in client.get("/api/samples").json()}
    assert "sales_2026" in samples
    doc_id = _process(client, "sales_2026")
    result = client.get(f"/api/process-result/{doc_id}").json()
    assert result["smart_stats"]["split_tables"] == 0
    assert result["smart_stats"]["type_counts"]["table"] > 0

    assert client.get(f"/api/history/{doc_id}").json()["counts"]["questions"] == 0
    assert client.get("/api/history").json() == []

    chat = client.post("/api/chat", json={"doc_id": doc_id, "method": "smart", "question": "Total revenue?"})
    assert chat.status_code == 200, chat.text
    assert chat.json()["prompt_sent"].startswith("You are a helpful assistant")

    calls_before = len(client.calls)  # type: ignore[attr-defined]
    history = [{"role": "user", "content": "Total revenue?"}, {"role": "assistant", "content": "27 million"}]
    compare = client.post("/api/compare", json={"doc_id": doc_id, "question": "and in total?", "history": history})
    assert compare.status_code == 200, compare.text
    body = compare.json()
    assert body["smart"]["standalone_question"] == body["naive"]["standalone_question"] == "What is the total revenue?"
    rewrites = [c for c in client.calls[calls_before:] if "Rewrite" in c[0]["content"]]  # type: ignore[attr-defined]
    assert len(rewrites) == 1  # compare rewrites once, not once per method

    listed = [schemas.DocHistory.model_validate(item) for item in client.get("/api/history").json()]
    assert len(listed) == 1 and listed[0].doc_id == doc_id
    assert listed[0].counts.questions == 2 and listed[0].counts.compare == 1 and listed[0].counts.smart == 1
    assert listed[0].recent[1].naive_answer is not None

    cleared = client.delete(f"/api/history/{doc_id}")
    assert cleared.json() == {"cleared": True, "doc_id": doc_id}
    assert client.get("/api/history").json() == []

    other = client.post("/api/auth/register", json={"username": "omar", "password": "secret123"}).json()["token"]
    denied = client.get(f"/api/history/{doc_id}", headers={"Authorization": f"Bearer {other}"})
    assert denied.status_code == 404 and denied.json()["error"]["code"] == "DOC_NOT_FOUND"


def test_xlsx_upload_and_bad_type(client: TestClient, tmp_path: Path) -> None:
    workbook = openpyxl.Workbook()
    workbook.active.append(["Item", "Qty"])
    workbook.active.append(["Pen", 3])
    path = tmp_path / "stock.xlsx"
    workbook.save(path)
    upload = client.post("/api/upload", files={"file": ("stock.xlsx", path.read_bytes(), "application/octet-stream")})
    assert upload.status_code == 200, upload.text
    assert upload.json()["pages"] == 1
    broken = client.post("/api/upload", files={"file": ("bad.xlsx", b"nope", "application/octet-stream")})
    assert broken.status_code == 422
    wrong = client.post("/api/upload", files={"file": ("slides.pptx", b"x", "application/octet-stream")})
    assert wrong.status_code == 415 and "Excel" in wrong.json()["error"]["message"]


def test_deleting_document_clears_history(client: TestClient) -> None:
    doc_id = _process(client, "company_handbook")
    assert client.post("/api/chat", json={"doc_id": doc_id, "method": "naive", "question": "Leave?"}).status_code == 200
    assert len(client.get("/api/history").json()) == 1
    assert client.delete(f"/api/documents/{doc_id}").status_code == 200
    assert client.get("/api/history").json() == []
